"""Generate the ResNet-results section for the thesis-defense deck.

Builds a small, self-contained set of PowerPoint slides covering ONLY the
ResNet findings from "Stabilizing Gradient Inversion in Federated Learning".
The slides are designed to be dropped into the group's existing deck, so the
styling is intentionally clean and neutral (white background, DTU-red accent).

Run:  python slides/make_resnet_slides.py
Out:  slides/resnet_results.pptx
"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

# ---------------------------------------------------------------- palette ----
DTU_RED = RGBColor(0x99, 0x00, 0x00)
INK = RGBColor(0x22, 0x22, 0x22)
SLATE = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x1B, 0x7A, 0x3D)
FONT = "Calibri"

ASSETS = Path(__file__).resolve().parent / "assets"
OUT = Path(__file__).resolve().parent / "resnet_results.pptx"

# 16:9 widescreen
EMU_W = Inches(13.333)
EMU_H = Inches(7.5)


def _set_bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def _box(slide, x, y, w, h, fill=None, line=None, line_w=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w or Pt(1)
    return shp


def _text(slide, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          space_after=6, line_spacing=1.0):
    """runs: list of paragraphs; each paragraph is a list of (text, size, bold,
    color, italic) run tuples, or a single such tuple."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        p.space_before = Pt(0)
        p.line_spacing = line_spacing
        if isinstance(para, tuple):
            para = [para]
        for (txt, size, bold, color, *rest) in para:
            italic = rest[0] if rest else False
            r = p.add_run()
            r.text = txt
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.name = FONT
            r.font.color.rgb = color
    return tb


def _bullets(slide, x, y, w, h, items, size=17, gap=10, color=INK):
    """items: list of (level, text, bold, color?) tuples."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        lvl, txt, bold = item[0], item[1], item[2]
        c = item[3] if len(item) > 3 else color
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.space_before = Pt(0)
        p.line_spacing = 1.05
        p.level = lvl
        bullet = "—  " if lvl == 0 else "•  "
        r = p.add_run()
        r.text = bullet + txt
        r.font.size = Pt(size - lvl * 2)
        r.font.bold = bold
        r.font.name = FONT
        r.font.color.rgb = c
    return tb


def _img_fit(slide, path, x, y, max_w, max_h, align="center", valign="middle"):
    """Place image scaled to fit inside (max_w, max_h) preserving aspect."""
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(max_w / iw, max_h / ih)
    w = Emu(int(iw * scale))
    h = Emu(int(ih * scale))
    if align == "center":
        x = x + (max_w - w) // 2
    elif align == "right":
        x = x + (max_w - w)
    if valign == "middle":
        y = y + (max_h - h) // 2
    elif valign == "bottom":
        y = y + (max_h - h)
    slide.shapes.add_picture(str(path), x, y, w, h)
    return x, y, w, h


def _header(slide, kicker, title):
    # top accent bar
    _box(slide, 0, 0, EMU_W, Inches(0.12), fill=DTU_RED)
    _text(slide, Inches(0.6), Inches(0.32), Inches(12.1), Inches(0.35),
          [(kicker.upper(), 13, True, DTU_RED)])
    _text(slide, Inches(0.6), Inches(0.62), Inches(12.1), Inches(0.8),
          [(title, 30, True, INK)])
    _box(slide, Inches(0.62), Inches(1.42), Inches(1.1), Pt(3), fill=DTU_RED)


def _footer(slide, page):
    _text(slide, Inches(0.6), Inches(7.02), Inches(9.0), Inches(0.35),
          [("Stabilizing Gradient Inversion in Federated Learning · ResNet results",
            10, False, SLATE)])
    _text(slide, Inches(11.6), Inches(7.02), Inches(1.2), Inches(0.35),
          [(page, 10, False, SLATE)], align=PP_ALIGN.RIGHT)


def blank_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _set_bg(s, WHITE)
    return s


# ============================================================== build deck ===
prs = Presentation()
prs.slide_width = EMU_W
prs.slide_height = EMU_H

# ---- Slide 1 : section divider ---------------------------------------------
s = blank_slide(prs)
_set_bg(s, DTU_RED)
_box(s, 0, Inches(2.55), EMU_W, Inches(2.4), fill=RGBColor(0x7A, 0x00, 0x00))
_text(s, Inches(0.9), Inches(2.0), Inches(11.5), Inches(0.5),
      [("RESULTS", 16, True, RGBColor(0xF2, 0xC0, 0xC0))])
_text(s, Inches(0.9), Inches(2.5), Inches(11.5), Inches(1.3),
      [("ResNet", 60, True, WHITE)])
_text(s, Inches(0.9), Inches(3.85), Inches(11.5), Inches(1.0),
      [("Does the masking benefit seen for VGG transfer to residual "
        "architectures?", 22, False, RGBColor(0xF6, 0xDD, 0xDD))], line_spacing=1.1)
_text(s, Inches(0.9), Inches(6.5), Inches(11.5), Inches(0.5),
      [("ResNet-18 (basic blocks) · ResNet-50 (bottleneck) · CIFAR-100 "
        "· iDLG baseline vs. layer-wise masking", 13, False,
        RGBColor(0xE8, 0xC8, 0xC8))])

# ---- Slide 2 : layer-wise ablation -----------------------------------------
s = blank_slide(prs)
_header(s, "ResNet — layer-wise masking", "Masking does not help — the inverse of VGG")
_img_fit(s, ASSETS / "resnet18_ablation.png",
         Inches(6.55), Inches(1.65), Inches(6.4), Inches(5.0))
_text(s, Inches(6.55), Inches(6.55), Inches(6.4), Inches(0.4),
      [("ResNet-18, CIFAR-100: ΔPSNR (masked − baseline). "
        "Left of the dashed line = worse.", 10, False, SLATE)],
      align=PP_ALIGN.CENTER)
_bullets(s, Inches(0.6), Inches(1.75), Inches(5.8), Inches(5.0), [
    (0, "Leave-one-layer-out, all four optimizer / pretraining settings.", True),
    (1, "Non-pretrained: no single-layer mask improves reconstruction — "
        "every clear effect is a degradation.", False),
    (0, "conv1 is the most damaging layer to mask.", True),
    (1, "L-BFGS: PSNR 17.53 → 15.23 dB (−2.3 dB), SSIM "
        "0.54 → 0.32 (−41%).", False),
    (1, "bn1 stays near baseline; batch-norm is neutral at batch size 1.", False),
    (0, "Pretraining barely shifts the pattern.", True),
    (1, "conv1 drop deepens to −3.1 dB (17.46 → 14.41 dB).", False),
    (1, "Only layer4 under Signed AdamW is marginally positive "
        "(15.05 → 15.43 dB) — a lone exception.", False),
    (0, "ResNet-50 reproduces the same qualitative pattern.", True, DTU_RED),
], size=16, gap=9)
_footer(s, "2")

# ---- Slide 3 : sweep + restarts -> architectural ---------------------------
s = blank_slide(prs)
_header(s, "ResNet — robustness checks", "The absence of benefit is architectural")
# two stacked figures on the right
_img_fit(s, ASSETS / "masking_sweep.png",
         Inches(6.5), Inches(1.7), Inches(6.45), Inches(2.45))
_img_fit(s, ASSETS / "resnet18_restarts_trend.png",
         Inches(6.5), Inches(4.25), Inches(6.45), Inches(2.55))
_bullets(s, Inches(0.6), Inches(1.8), Inches(5.7), Inches(5.0), [
    (0, "Masking-percentage sweep (top).", True),
    (1, "ResNet-18 stays flat and trends down — peaks at only "
        "15 / 100 images near 40% masking.", False),
    (1, "Contrast: VGG-13 climbs to 40 / 100 at 70% masking.", False),
    (0, "Random restarts (bottom).", True),
    (1, "k = 1 → 5 restarts; keep the lowest-loss run.", False),
    (1, "Masked run never overtakes the unmasked baseline "
        "(gap 0.09 → 0.22 dB).", False),
    (0, "So it is not a bad-initialization artifact — even with "
        "equal restart budget, masking does not win.", True, DTU_RED),
], size=16, gap=9)
_footer(s, "3")

# ---- Slide 4 : why + takeaway ----------------------------------------------
s = blank_slide(prs)
_header(s, "ResNet — interpretation", "Why residual networks resist masking")
_bullets(s, Inches(0.6), Inches(1.8), Inches(12.1), Inches(3.4), [
    (0, "The information is present — not missing.", True, GREEN),
    (1, "ResNet-18 and ResNet-50 reach full column rank (d = 3,072) under "
        "both entry-selection modes → gradients still encode enough "
        "independent constraints on the input.", False),
    (0, "Skip connections distribute gradient information.", True),
    (1, "Earlier activations re-enter later blocks, spreading input-relevant "
        "signal across many residual paths. Removing a whole layer group "
        "discards useful constraints along with any poorly conditioned ones.", False),
    (0, "It is not parameter count.", True),
    (1, "On the convolutional backbone that constrains the pixels, the two are "
        "comparable (~9.4 M VGG-13 vs. ~11.5 M ResNet-18); ResNet just lacks "
        "VGG’s large maskable FC classifier.", False),
], size=16, gap=8)
# takeaway band
_box(s, Inches(0.6), Inches(5.45), Inches(12.13), Inches(1.25), fill=LIGHT)
_box(s, Inches(0.6), Inches(5.45), Inches(0.12), Inches(1.25), fill=DTU_RED)
_text(s, Inches(0.95), Inches(5.62), Inches(11.6), Inches(1.0),
      [[("Takeaway:  ", 18, True, DTU_RED),
        ("for ResNet, masking offers no consistent benefit — but the "
         "failure reflects the difficulty of isolating disruptive components "
         "in a residual architecture, ", 18, False, INK),
        ("not an absence of recoverable information.", 18, True, INK)]],
      anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.1)
_footer(s, "4")

prs.save(str(OUT))
print(f"Wrote {OUT}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
